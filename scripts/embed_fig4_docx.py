#!/usr/bin/env python3
"""把图 4（语义软代价层 Gazebo 闭环）嵌入权威载体 docs/语义地图.docx。

按 docx 现有图片格式（居中图片段落 + 后跟居中 caption 段落），
在 §4.7「语义软代价层 Gazebo 闭环」bullet 之后插入：
  1. 居中图片段落（fig4_semantic_cost_loop.png，宽 6.00in，与图 1/图 2 全宽图一致）
  2. 居中 caption 段落（图 4　语义软代价层 Gazebo 闭环轨迹…，Times New Roman 9pt bold）

用法（项目根 /home/saw/ros2_ws/src/semphyspc/ 下）：
    python3 scripts/embed_fig4_docx.py
"""
import copy
import shutil
import docx
from docx.text.paragraph import Paragraph
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

SRC = 'docs/语义地图.docx'
BAK = 'docs/语义地图.docx.bak-20261002-fig4'
IMG = 'paper/figures/fig4_semantic_cost_loop.png'
CAPTION = '图 4　语义软代价层 Gazebo 闭环轨迹（绕开致命行人、压过可穿越植被）'


def find_para(doc, needle, exclude=None):
    for p in doc.paragraphs:
        if needle in p.text:
            if exclude and any(e in p.text for e in exclude):
                continue
            return p
    raise RuntimeError(f'段落未找到: {needle}')


def main():
    shutil.copy2(SRC, BAK)
    doc = docx.Document(SRC)

    # 图 1 的图片段落与 caption 段落作为格式模板
    p_fig1_img = None
    p_fig1_cap = None
    for i, p in enumerate(doc.paragraphs):
        if p.text.startswith('图 1　') and p_fig1_cap is None:
            p_fig1_cap = p
            p_fig1_img = doc.paragraphs[i - 1]
    assert p_fig1_img is not None and p_fig1_cap is not None

    # 锚点：§4.7 语义软代价层 bullet（make_md_para 插入，含「语义软代价层 Gazebo 闭环」）
    anchor = find_para(doc, '语义软代价层 Gazebo 闭环')

    # 1) 图片段落：deepcopy 图1图片段落，清空 runs（含旧图），加图4
    img_el = copy.deepcopy(p_fig1_img._p)
    img_p = Paragraph(img_el, p_fig1_img._parent)
    for r in list(img_p.runs):
        r._r.getparent().remove(r._r)
    img_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = img_p.add_run()
    run.add_picture(IMG, width=Inches(6.0))

    # 2) caption 段落：deepcopy 图1 caption，保留格式，替换文本
    cap_el = copy.deepcopy(p_fig1_cap._p)
    cap_p = Paragraph(cap_el, p_fig1_cap._parent)
    for r in list(cap_p.runs):
        r._r.getparent().remove(r._r)
    cap_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cr = cap_p.add_run(CAPTION)
    cr.font.name = 'Times New Roman'
    cr.font.size = Pt(9)
    cr.font.bold = True
    # 中文回退宋体
    from docx.oxml.ns import qn
    cr._element.get_or_add_rPr()
    rf = cr._element.rPr.get_or_add_rFonts()
    rf.set(qn('w:eastAsia'), '宋体')

    # 插入：anchor 之后 -> 图片段落 -> caption 段落
    anchor._p.addnext(img_el)
    img_el.addnext(cap_el)

    doc.save(SRC)

    # 验证图片数 7 -> 8
    d2 = docx.Document(SRC)
    print(f'完成：{SRC} 已更新（备份 -> {BAK}）')
    print(f'图片数：{len(d2.inline_shapes)}（原 7）')
    print(f'最后一张图片宽 {d2.inline_shapes[-1].width/914400:.2f}in x '
          f'高 {d2.inline_shapes[-1].height/914400:.2f}in')


if __name__ == '__main__':
    main()
