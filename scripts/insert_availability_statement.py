#!/usr/bin/env python3
"""在结论之后、参考文献之前插入 Data/Code availability 声明（双版同步）。

英文：加粗标签 Data Availability / Code Availability，正文 10pt。
中文：对应中文标签，正文 10.5pt。

用法（项目根下）：python3 scripts/insert_availability_statement.py
"""
import shutil

import docx
from docx.oxml.ns import qn
from docx.shared import Pt

JOBS = [
    dict(
        path='docs/paper_draft_semphyspc.docx',
        ref_prefix='References',
        size=10,
        items=[
            ('Data Availability: ',
             'The nuScenes dataset is publicly available at https://www.nuscenes.org; the semantic-to-physical '
             'mapping library, trained model checkpoints, and supplementary results will be released via IEEE '
             'DataPort upon acceptance.'),
            ('Code Availability: ',
             'The source code and scripts will be released at https://github.com/[repository] upon acceptance.'),
        ],
    ),
    dict(
        path='docs/语义地图.docx',
        ref_prefix='参考文献',
        size=10.5,
        items=[
            ('数据可得性：',
             'nuScenes 数据集公开可得（https://www.nuscenes.org）；语义→物理材质映射库、训练模型权重与补充结果将于录用后经 '
             'IEEE DataPort 公开。'),
            ('代码可得性：',
             '源代码与脚本将于录用后发布于 https://github.com/[repository]。'),
        ],
    ),
]


def _mk_run(p, text, size, bold):
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(size)
    run.font.bold = bold
    rf = run._element.get_or_add_rPr().get_or_add_rFonts()
    rf.set(qn('w:eastAsia'), '宋体')
    return run


def main():
    for job in JOBS:
        shutil.copy2(job['path'], job['path'] + '.bak-availability')
        doc = docx.Document(job['path'])
        anchor = None
        for p in doc.paragraphs:
            if p.text.strip() == job['ref_prefix']:
                anchor = p
                break
        assert anchor is not None, f"{job['path']}: 未找到参考文献标题 {job['ref_prefix']}"

        for label, body in job['items']:
            p = anchor.insert_paragraph_before()
            p.paragraph_format.space_after = Pt(6)
            _mk_run(p, label, job['size'], bold=True)
            _mk_run(p, body, job['size'], bold=False)

        doc.save(job['path'])

        d2 = docx.Document(job['path'])
        hits = [p.text.strip()[:40] for p in d2.paragraphs if 'Availability' in p.text or '可得性' in p.text]
        print(f"{job['path']}: 声明已插入 {len(hits)} 段 -> {hits}")


if __name__ == '__main__':
    main()
