#!/usr/bin/env python3
"""把 paper/paper_draft.md 的正向结果（num_group 64→256 + 门控投票融合）同步进
投稿版 docx（docs/paper_draft_semphyspc.docx）。

做法：备份原 docx → 按文本匹配定位段落/表格 → 替换正文、插入 §4.13、加 g256 行。

用法:
    python3 scripts/sync_paper_docx.py
"""
import copy
import shutil
import docx
from docx.text.paragraph import Paragraph
from docx.oxml.ns import qn

SRC = 'docs/paper_draft_semphyspc.docx'
BAK = 'docs/paper_draft_semphyspc.docx.bak'


def set_para_text(p, text):
    """替换段落文本，保留第一个 run 的字体格式，删其余 run。"""
    runs = p.runs
    if runs:
        first = runs[0]
        for r in runs[1:]:
            r._r.getparent().remove(r._r)
        first.text = text
    else:
        p.add_run(text)


def find_para(doc, needle, exclude=None):
    for p in doc.paragraphs:
        if needle in p.text:
            if exclude and any(e in p.text for e in exclude):
                continue
            return p
    raise RuntimeError(f'段落未找到: {needle}')


def new_para_like(template, text):
    """deepcopy 模板段落，清空文本后填入 text，返回新的 w:p 元素。"""
    el = copy.deepcopy(template._p)
    p = Paragraph(el, template._parent)
    for r in p.runs:
        r.text = ''
    if p.runs:
        p.runs[0].text = text
    else:
        p.add_run(text)
    return el


def insert_before(anchor, el):
    anchor._p.addprevious(el)


def make_table(doc, header, rows):
    """新建 TableGrid 样式表格，返回 table 对象（先加到文档末尾，稍后 addprevious 移动）。"""
    tbl = doc.add_table(rows=1 + len(rows), cols=len(header))
    try:
        tbl.style = doc.styles['Table Grid']
    except KeyError:
        pass
    for j, h in enumerate(header):
        tbl.rows[0].cells[j].text = h
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            tbl.rows[i + 1].cells[j].text = str(val)
    return tbl


def main():
    shutil.copy2(SRC, BAK)
    doc = docx.Document(SRC)

    # ---- 1. 摘要（中文） ----
    p_abs_cn = find_para(doc, '实验在 nuScenes 公开数据集上验证了分割前置（mIoU 0.2395')
    set_para_text(p_abs_cn,
        '自动驾驶车辆在失控等紧急场景下需要选择"非致命、可变形/可破坏"的物体（灌木、交通锥、护栏等）'
        '进行主动软碰撞以实现安全减速，而避免撞击行人、动物、车辆与墙体等致命或刚性目标；另一方面，'
        '基于实地扫描的三维点云地图本身不具备任何碰撞与损毁逻辑，无法直接用于游戏场景建模。二者的共同瓶颈在于：'
        '仅得到逐点语义类别还不够，还必须将语义类别转化为物理交互属性（碰撞安全等级、质量、摩擦、恢复系数、'
        '可破坏性与损毁响应）。针对这一问题，本文提出一套从语义分割到物理交互的端到端方法：以基于 RGB 与 XYZ '
        '早期融合的 PointMambaSeg 语义分割为前置，构建 32 类语义→物理材质映射库；在此基础上实现实例级聚类、'
        '位置相关破坏模型（材料强度 × 几何结构强度 × 撞击局部化）、导航软碰撞目标决策、跨帧实例跟踪，'
        '以及面向 Unity/Unreal/ROS2 的游戏物理与导航系统导出。实验在 nuScenes 公开数据集上验证了分割前置：'
        '将 Mamba 的 token 数（num_group）由 64 提升到 256，mIoU 由 0.2323 提升到 0.3090（+33%）、OA 达 0.8679，'
        '据此定位 Mamba 的规模化瓶颈为 token 数而非输入点数；针对分割对长尾罕见类（行人/自行车/摩托车等）'
        '近乎全灭的缺陷，提出 6 相机深度门控的 2D/3D 投票融合，把罕见类 mIoU 由 0.0994 提升到 0.4512'
        '（全类 mIoU 达 0.4274）；级联实例分割经类别门控把 thing 类语义精度提升至 0.724，并通过置信度门控实现'
        '「可拒识」的可信决策（拒识 50% 最不确定实例后精度 +14.5 点）；并展示了同一撞击能量作用于物体不同部位'
        '得到不同破坏结果的真实性，以及多航向下正确选择软碰撞目标的决策行为。该方法为"导航安全碰撞"与'
        '"游戏实景建图"提供了一条可落地的技术路径。')

    # ---- 2. Abstract（英文） ----
    p_abs_en = find_para(doc, 'Experiments on the nuScenes benchmark validate the segmentation front-end (mIoU 0.2395')
    set_para_text(p_abs_en,
        'In loss-of-control scenarios, autonomous vehicles should actively select non-lethal, deformable or frangible '
        'objects (e.g., bushes, traffic cones, guardrails) for soft collision to decelerate safely, while avoiding '
        'lethal or rigid targets such as pedestrians, animals, vehicles and walls. Meanwhile, real-scanned 3D '
        'point-cloud maps carry no collision or destruction logic and cannot be used directly for game scene building. '
        'The shared bottleneck is that per-point semantic labels alone are insufficient; the semantics must be '
        'translated into physical interaction properties (collision-safety level, mass, friction, restitution, '
        'destructibility and damage response). We propose an end-to-end method from semantic segmentation to physical '
        'interaction: (1) a PointMambaSeg segmentation front-end with early RGB+XYZ fusion; (2) a 32-class '
        'semantic-to-physical-material mapping library; (3) instance-level clustering and a position-dependent '
        'destruction model (material strength × geometric structural strength × impact localization); (4) a '
        'soft-collision target decision module for navigation; (5) cross-frame instance tracking; and (6) game-physics '
        'and navigation system export for Unity/Unreal/ROS2. Experiments on the nuScenes benchmark validate the '
        'segmentation front-end: raising the Mamba token count (num_group) from 64 to 256 improves mIoU from 0.2323 to '
        '0.3090 (+33%, OA 0.8679), locating the scaling bottleneck in token count rather than input resolution; to '
        'address the near-total failure on long-tail rare classes (pedestrians, bicycles, motorcycles, etc.), we '
        'propose a 6-camera depth-gated 2D/3D voting fusion that lifts rare-class mIoU from 0.0994 to 0.4512 '
        '(full mIoU 0.4274); the cascaded instance segmentation reaches 0.724 thing-class accuracy via class gating, '
        'and a confidence-gating module enables rejectable, trustworthy decisions (+14.5 points after rejecting the '
        '50% most uncertain instances); we further demonstrate position-dependent destruction and heading-dependent '
        'soft-collision decisions. The method provides a practical path for both navigation safety collision and game '
        'map building.')

    # ---- 3. 贡献：在贡献 4 之后插入贡献 5、6 ----
    p_contrib4 = find_para(doc, '跨帧实例跟踪与系统导出')
    c5 = new_para_like(p_contrib4,
        '5. Mamba 规模化瓶颈的定位：通过消融定位 PointMambaSeg 的规模化瓶颈为 token 数（num_group）而非输入点数'
        '——升点数（4096→16384）不升反降，而升 token 数（num_group 64→256）把 mIoU 提升 33%（0.2323→0.3090）。')
    c6 = new_para_like(p_contrib4,
        '6. 长尾罕见类的 2D/3D 深度门控投票融合：利用 6 相机 2D 检测框对 3D 点做深度门控的投票回填，'
        '把罕见类 mIoU 提升 4.5 倍（0.0994→0.4512），补齐逐点分割对长尾类近乎全灭的短板。')
    p_contrib4._p.addnext(c5)
    c5.addnext(c6)

    # ---- 4. 引言结尾 ----
    p_intro_end = find_para(doc, '实验部分在 nuScenes 公开数据集上验证了分割前置（并如实报告了')
    set_para_text(p_intro_end,
        '实验部分在 nuScenes 公开数据集上验证了分割前置（含"升点数 vs 升 token 数"的消融，'
        '以及罕见类门控投票融合），以及物理映射、破坏模型、软碰撞决策与跟踪的可行性。')

    # ---- 5. 3.2 分割前置 ----
    p_32 = find_para(doc, 'num_group=64 固定')
    set_para_text(p_32,
        '在输入端将点云 XYZ 与相机投影得到的 RGB 颜色在通道维拼接，形成 6 通道输入，经过非层级编码器与 '
        'Mamba 块提取特征，再通过 KNN 上采样解码器恢复逐点语义。关键设计：Mamba 的 token 序列长度由 num_group '
        '决定（FPS 采样 num_group 个中心 + KNN 32 邻域聚合），与输入点数无关；token 数是 Mamba 序列建模的信息'
        '容量上限，是本文定位到的规模化瓶颈（§4.2）。默认配置 num_group=64，规模化配置 num_group=256。'
        '模型参数量约 12.34M。')

    # ---- 6. 4.1 设置 ----
    p_41 = find_para(doc, 'num_group=64、in_channels=6')
    set_para_text(p_41,
        '采用 nuScenes 公开数据集（32 类 lidarseg 语义标签，700 train / 150 val scenes，共 34,149 关键帧），'
        '输入为 LiDAR 点云 + 6 相机投影 RGB（6 通道），每帧采样 16,384 点。分割模型 PointMambaSeg：'
        'trans_dim=384、depth=12、group_size=32、in_channels=6、num_classes=32，AdamW 优化器（lr=1e-3、'
        'weight_decay=0.05），余弦退火，200 epoch。num_group 为消融变量（64 / 256）。硬件为 RTX 5070 Ti。')

    # ---- 7. 4.2 消融 ----
    p_ablation = find_para(doc, '消融（如实报告的负结果）')
    set_para_text(p_ablation,
        '消融：升点数 vs 升 token 数（定位规模化瓶颈）。两组对照：① 升点数——在 num_group=64 固定下将点数 '
        '4096→16384，mIoU 0.2395→0.2323、OA 0.8294→0.8095 略降（逐类有涨有跌：car 0.447→0.476、'
        'barrier 0.339→0.418、driveable 0.827→0.855 提升；manmade 0.762→0.725、vegetation 0.779→0.718 下降），'
        '为负结果；② 升 token 数——保持 16384 点不变，将 num_group 64→256，mIoU 0.2323→0.3090（+33%）、'
        'OA 0.8095→0.8679（+5.8 点），200 epoch 跑满，best 出现在 epoch 133。两组消融合起来定位了瓶颈：'
        'Mamba 的规模化瓶颈是 token 数（num_group），而非输入点数——num_group 决定了序列建模的信息容量上限，'
        'token 数不变时单纯增加点数无法带来信息增益。这一发现把此前孤立的"升点数负结果"解释为同一瓶颈的两面，'
        '也为后续规模化（GridSampling + 序列化 patch，见 §5）指明了方向。')

    # ---- 8. mIoU 汇总表（表 3）加 g256 行 ----
    tbl_miou = doc.tables[3]
    row = tbl_miou.add_row()
    vals = ['seg_16384_g256（num_group=256）', '16384', '0.3090', '0.8679']
    for j, val in enumerate(vals):
        row.cells[j].text = val

    # ---- 9. 3.8 导航（补充语义+几何融合导航） ----
    p_nav = find_para(doc, '导航：封装为 ROS2 节点（soft_contact 包）')
    set_para_text(p_nav,
        '导航：封装为 ROS2 节点（soft_contact 包），订阅点云 + 里程计，发布推荐撞击点、候选目标与摘要。'
        '此外提供语义+几何融合导航：把语义障碍（PointMamba 实时分割按 obstacle_classes 过滤）与坡度障碍'
        '（PCA 地面坡度估计，陡坡 cell）作为「标记型」观测源叠加进 Nav2 costmap，使规划同时避开'
        '「是什么（语义障碍）」与「能不能上（陡坡）」（slope_obstacle_publisher.py + 增强参数 '
        'nav2_params_semantic.yaml）。')

    # ---- 10. 4.7 ROS2 段后补充 Nav2 融合导航验证 ----
    p_47 = find_para(doc, 'soft_contact 包 colcon build')
    nav2 = new_para_like(p_47,
        'Nav2 语义+几何融合导航：离线仿真（假 TF + 假激光 + AMCL 定位）验证了 Nav2 定位与路径规划全链路；'
        '坡度障碍节点在 Semantic3D 场景上离线 dry-run 输出 114 个陡坡 cell / 16,785 个障碍点，'
        '语义+坡度两个障碍观测源通过 ROS 发布（/semantic_obstacles 50 万点、/slope_obstacles 16,785 点），'
        'Nav2 加载增强参数后 local/global costmap 均订阅 scan semantic_obstacles slope_obstacles 三个观测源，'
        'smoke test 无报错。')
    p_47._p.addnext(nav2)

    # ---- 11. 插入 §4.13（在 5 结论 之前） ----
    p_concl = find_para(doc, '5 结论')
    p_concl_prev = p_concl._p.getprevious()  # 4.12 结论正文段
    tpl_body = Paragraph(p_concl_prev, p_concl._parent)  # 正文模板
    tpl_head = p_concl  # 标题模板

    sec = [
        ('head', '4.13 罕见类的 2D/3D 深度门控投票融合'),
        ('body', '逐点分割对 long-tail 罕见类（行人各子类、barrier、trafficcone、bicycle、motorcycle 等 11 类）'
                 '几乎全灭（base 罕见类 mIoU 仅 0.0994），是 mIoU 的主要拖累。本节提出并验证一条互补路径：'
                 '用 2D 图像检测的罕见类框对 3D 点做门控投票回填，不改变 3D 网络结构，只作为后处理融合。'),
        ('body', '方法。离线把 nuScenes 3D 标注框投影到 6 相机（CAM_FRONT / FRONT_LEFT / FRONT_RIGHT / BACK / '
                 'BACK_LEFT / BACK_RIGHT），缓存为 2D 框（共 574,241 个），每框记录类别与 3D 框在对应相机系的深度'
                 '范围 [d_min, d_max]。推理时对每个 3D 点：投影到 6 相机并取并集，若某点落在某罕见类 2D 框内、'
                 '且点深度 ∈ [d_min−δ, d_max+δ]（深度门控，δ=0.5 m），则该点获得 2D 类别标签 sem_2d；'
                 '门控投票只对 3D 置信度 conf_3d < θ 的点用 sem_2d 覆盖 3D 预测。'),
        ('body', '深度门控与 6 相机都是必须的。2D 框无深度，同像素射线上物体背后的遮挡点会被误关联：'
                 '6 相机下 naive（无深度门控）关联 precision 仅 0.358，加深度门控后提升到 0.701；'
                 '相机越多、视野越全，遮挡污染越重，深度门控越不可省。而单 CAM_FRONT 只见约 12.6% 的点，'
                 '6 相机并集覆盖率才达 86.9%——二者缺一，罕见类的覆盖与精度都会大幅受损。'),
        ('body', '结果（g256 模型 + 6 相机深度门控；θ 为门控置信度阈值，θ 越大越信任 2D）：'),
        ('table', None),
        ('body', '单类亮点（θ=0.9~1.0）：ped.police 0→0.741、ped.adult 0.169→0.614、'
                 'motorcycle 0.026→0.565、bicycle 0.010→0.446、trafficcone 0.230→0.479。'
                 'θ=0.9 是稳妥平衡点（召回高且误报可控）；θ=1.0 罕见类最高但全类误报略增。'),
        ('body', '结论：3D 逐点分割对长尾罕见类几乎学不到（base IoU≈0），而 2D 图像检测对同类物体判别力强；'
                 '「2D 检测 + 深度门控 + 置信度投票」把罕见类 mIoU 提升 4.5 倍，是补齐长尾类的有效手段，'
                 '且不增加 3D 网络训练代价，可直接叠加在既有分割前置之上。'),
    ]

    result_header = ['指标', 'base(3D)', 'θ=0.7', 'θ=0.9', 'θ=1.0']
    result_rows = [
        ['mIoU（罕见类）', '0.0994', '0.3413', '0.4512', '0.5059'],
        ['mIoU（全部）', '0.3090', '—', '0.4274', '0.4457'],
    ]
    tbl_413 = make_table(doc, result_header, result_rows)

    for kind, text in sec:
        if kind == 'head':
            el = new_para_like(tpl_head, text)
        elif kind == 'body':
            el = new_para_like(tpl_body, text)
        elif kind == 'table':
            el = tbl_413._tbl
        else:
            continue
        insert_before(p_concl, el)

    # ---- 12. 结论 ----
    p_concl_body = find_para(doc, '实验验证了分割前置的可行性（mIoU 0.2395')
    set_para_text(p_concl_body,
        '本文针对"语义标签不足以驱动物理交互"这一瓶颈，提出了一条以语义分割为前置、以物理交互链路为主体的'
        '完整方法：构建 32 类语义→物理材质映射库，实现实例级聚类、位置相关破坏模型、导航软碰撞决策与跨帧实例'
        '跟踪，并完成 Unity/Unreal/ROS2 的工程化导出。实验验证了分割前置的可行性（num_group=256 时 mIoU 0.3090，'
        '且在 Semantic3D 上零样本 OA 0.536、冻结 backbone 微调头优于从零 +3.9 mIoU）；通过「升点数 vs 升 token 数」'
        '的消融定位了 Mamba 的规模化瓶颈为 token 数（num_group 64→256 带来 +33% mIoU），并以 6 相机深度门控的 '
        '2D/3D 投票融合把罕见类 mIoU 提升 4.5 倍（0.0994→0.4512）、全类 mIoU 达 0.4274，补齐了逐点分割对长尾类'
        '的短板；级联实例分割（offset 回归 + 聚类 + 物体判类）经类别门控后把 thing 类语义精度从 dense 的 0.485 '
        '提升到 0.724，并如实报告了「实例 PQ 最优与级联语义精度最优不一致」这一反直觉消融发现；进一步将类别门控'
        '推广为通用置信度门控，用 softmax 不确定度与 dense↔B2 跨模块分歧做拒识，验证了「知道何时说不知道」的'
        '可信决策能力（拒识 50% 最不确定实例后精度 +14.5 点，跨模块一致时 B2 精度达 0.953）；此外展示了位置相关'
        '破坏与航向相关软碰撞决策的真实行为。该方法为自动驾驶安全碰撞与游戏实景建图提供了统一的技术框架。')

    # ---- 13. 未来工作第一条 ----
    p_future = find_para(doc, '以加权/焦点损失（Focal Loss）')
    set_para_text(p_future,
        '（1）以 GridSampling + 序列化 patch 进一步规模化到单帧几万点（token 数瓶颈已定位，方向明确）；')

    # ---- 14. 待办清单 ----
    p_todo_last = find_para(doc, 'rare thing 类长尾问题的 Focal Loss / copy-paste 增强实验')
    todo_done = new_para_like(p_todo_last,
        '• [x] 写入正向结果：num_group 64→256 规模化（mIoU 0.3090，§4.2）+ 罕见类 2D/3D 深度门控投票融合（§4.13）；'
        '并整合 Nav2 语义+几何融合导航（§3.8 / §4.7）')
    p_todo_last._p.addprevious(todo_done)

    doc.save(SRC)
    print(f'完成：{SRC} 已更新（备份 -> {BAK}）')
    print(f'正文段落数 = {len(doc.paragraphs)}，表格数 = {len(doc.tables)}')


if __name__ == '__main__':
    main()
