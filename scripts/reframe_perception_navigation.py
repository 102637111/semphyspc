#!/usr/bin/env python3
"""阶段 A 定位重写：把论文定位从「感知 + 碰撞」重构为「感知 + 导航」，物理碰撞降为辅助分析。

改动范围（权威载体 docs/语义地图.docx）：
  - 标题（中/英）：去「碰撞」、以「导航安全」收尾
  - 摘要（中/英）：以「导航安全感知」开篇，方法框定「语义分割为核心、物理属性建模为辅助」
  - 关键词（中/英）：导航/语义代价地图提前，碰撞/材质后移
  - 贡献：以语义分割（感知）+ 语义软代价导航为主线，物理属性建模与撞击标定为辅助
  - 结论：同一重构措辞

运行前自动备份到 /tmp/语义地图.docx.bak-reframe-<ts>，可回滚。
用法:
    python3 scripts/reframe_perception_navigation.py [--docx docs/语义地图.docx]
"""
import argparse
import shutil
import time

import docx


def backup(path):
    ts = time.strftime('%Y%m%d-%H%M%S')
    bak = f'/tmp/{path.split("/")[-1]}.bak-reframe-{ts}'
    shutil.copy2(path, bak)
    return bak


def set_para_text(p, text):
    """整段替换文本，保留第一 run 的字体（bold/italic/size/name）。"""
    if p.runs:
        src = p.runs[0]
        bold = src.bold
        italic = src.italic
        size = src.font.size
        name = src.font.name
    else:
        bold = italic = None
        size = name = None
    # 清空所有 run
    for r in list(p.runs):
        r._element.getparent().remove(r._element)
    run = p.add_run(text)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if size is not None:
        run.font.size = size
    if name is not None:
        run.font.name = name


def replace_substr(p, old, new):
    """在段落全文内做子串替换，返回替换次数；格式坍缩为单 run（可接受）。"""
    text = p.text
    n = text.count(old)
    if n == 0:
        return 0
    set_para_text(p, text.replace(old, new))
    return n


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--docx', default='docs/语义地图.docx')
    args = ap.parse_args()

    bak = backup(args.docx)
    print(f'已备份到 {bak}')

    d = docx.Document(args.docx)
    paras = d.paragraphs

    # ---- 全文替换（标题/摘要/关键词/贡献）----
    full = {
        0: '面向导航安全的语义点云分割与物理属性建模方法',
        1: 'Semantic Point-Cloud Segmentation and Physical-Property Modeling for Autonomous Navigation Safety',
        4: ('自动驾驶在复杂环境中需要语义级的环境理解来支撑安全导航：不仅要识别“前方是什么”'
            '（逐点语义分割），还要把语义转化为导航系统可直接消费的代价信息，并据此在紧急场景下做出安全决策。'
            '针对这一问题，本文提出一套以语义点云分割为核心、以物理属性建模为辅助的方法：'
            '基于 RGB+XYZ 早期融合的 PointMambaSeg 分割前置、32 类语义→物理材质映射库、撞击后果标定、'
            '软碰撞决策模块与面向 ROS2 的导航系统导出。在 nuScenes 数据集上，分割前置达到 0.3923 mIoU，'
            '严格配对消融表明 RGB 早期融合贡献 +0.0331 mIoU；针对长尾分布下罕见致命类（行人、动物）近乎全灭的问题，'
            '引入 2D 检测与深度门控投票，把致命类被误判为可撞类的 critical_rate 从 5.30% 压至 3.80%'
            '（GT oracle 上界 0.92%）；软碰撞决策落地为语义软代价层，并在 Gazebo 物理仿真闭环中验证了'
            '规划器仅凭连续软代价即可绕开致命行人、穿越可穿越植被。该方法为自动驾驶导航安全感知提供了一条可落地的技术路径。'),
        5: '关键词：点云语义分割；Mamba；导航安全；语义代价地图；致命类安全；物理材质映射',
        8: 'Keywords: point-cloud semantic segmentation; Mamba; navigation safety; semantic cost map; '
           'lethal-class safety; physical material mapping',
        14: ('为此，本文提出一条以语义点云分割为核心、以导航安全为主线、以物理属性建模为辅助的完整方法，'
             '命名为 SemPhysPC（Semantic Physical Point Cloud，语义物理点云框架），主要贡献如下：'),
        15: ('1. 语义分割前置（感知核心）：基于 RGB+XYZ 早期融合的 PointMambaSeg 逐点分割，'
             '在 nuScenes 上达 0.3923 mIoU（OA 0.9027）；通过「升点数 vs 升 token 数」消融定位规模化瓶颈为 token 数，'
             'RGB 严格配对消融净贡献 +0.0331 mIoU。'),
        16: ('2. 语义软代价层与导航集成（导航主线）：把逐点语义转化为规划器可消费的连续软代价'
             '（致命类 254、可穿越类 30~80），接入 ROS2/Nav2 导航栈，在 Gazebo 闭环与 Semantic3D 真实地图上'
             '验证规划器绕开致命行人、穿越可穿越植被。'),
        17: ('3. 物理属性建模与撞击标定（辅助分析）：将 nuScenes 32 个语义类别映射到 5 级碰撞安全等级'
             '与 13 类物理材质（密度、摩擦、恢复系数、可破坏性、致命性代价），并以闭合形式解标定各类撞击的减速量 Δv，'
             '定量指出原软碰撞候选集的减速上限仅初速的 2.75%，据此把决策扩展为「正碰软目标 + 刮蹭硬面」双模式。'),
        18: ('4. 致命类安全与跨帧跟踪：基于匈牙利匹配实现跨帧稳定实例 ID 与静态点云融合；'
             '引入 2D 检测 + 深度门控投票，把致命类被误判为可撞类的 critical_rate 从 5.30% 压至 3.80%。'),
        19: ('实验部分在 nuScenes 公开数据集上验证了分割前置与语义软代价导航'
             '（并如实报告了“升点数”消融的负结果），物理映射、撞击标定、软碰撞决策与跟踪作为辅助分析一并验证。'),
    }
    for idx, text in full.items():
        set_para_text(paras[idx], text)
        print(f'[全文替换] 段落 {idx} 完成')

    # ---- 子串替换（英文摘要 & 结论：只动框架句，保留中间大量数字）----
    subs = {
        7: [
            ('In loss-of-control scenarios, autonomous vehicles should actively select non-lethal, '
             'deformable or frangible objects (e.g., bushes, traffic cones, guardrails) for soft collision '
             'to decelerate safely, while avoiding lethal or rigid targets such as pedestrians, animals, '
             'vehicles and walls.',
             'Autonomous driving in complex environments demands semantic-level scene understanding to support '
             'safe navigation: beyond recognizing what lies ahead (per-point semantic segmentation), the semantics '
             'must be translated into cost information that a planner can directly consume, so that safe decisions '
             'can be made even in emergency scenarios.'),
            ('The bottleneck is that per-point semantic labels alone are insufficient; the semantics must be '
             'translated into physical interaction properties (collision-safety level, mass, friction, restitution, '
             'destructibility). We propose an end-to-end method from semantic segmentation to physical interaction:',
             'The bottleneck is that per-point semantic labels alone are insufficient for navigation; they must be '
             'translated into navigation-consumable cost (collision-safety level and traversability), with '
             'physical-material properties (mass, friction, restitution, destructibility) as an auxiliary analysis. '
             'We propose a framework centered on semantic point-cloud segmentation with physical-property modeling '
             'as auxiliary:'),
            ('The method provides a practical path for navigation safety collision (and is extensible to game '
             'map building).',
             'The method provides a practical path for navigation-safe perception in autonomous driving (and is '
             'extensible to game map building).'),
        ],
        134: [
            ('本文针对“语义标签不足以驱动物理交互”这一瓶颈',
             '本文针对“逐点语义标签不足以直接支撑安全导航”这一瓶颈'),
            ('以语义分割为前置、以物理交互链路为主体',
             '以语义点云分割为核心、以导航安全为主线、以物理属性建模为辅助'),
            ('该方法为自动驾驶导航安全碰撞提供了统一的技术框架',
             '该方法为自动驾驶导航安全感知提供了统一的技术框架'),
        ],
    }
    for idx, repls in subs.items():
        for old, new in repls:
            n = replace_substr(paras[idx], old, new)
            if n != 1:
                print(f'[警告] 段落 {idx} 子串替换命中 {n} 次（期望 1）：{old[:40]}...')
            else:
                print(f'[子串替换] 段落 {idx} 命中 1 次：{old[:40]}...')

    d.save(args.docx)
    print(f'已保存 {args.docx}')


if __name__ == '__main__':
    main()
