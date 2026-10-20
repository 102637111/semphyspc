#!/usr/bin/env python3
"""把图 4/图 5 合并成一张三子图，写回权威载体 docs/语义地图.docx。

操作（§4.6 系统集成验证区域，按 body 子元素索引定位）：
  1. 删除原「图 5 图片段落 + 图 5 caption 段落」。
  2. 把「真实地图语义导航验证」文字段落移到「图 4 图片段落」之前，
     使顺序变为：Gazebo 闭环文字 → 真实地图文字 → 三子图 → caption。
  3. 把「图 4 图片段落」内的图片替换为 fig4_system_integration.png（宽 6.00in）。
  4. 把「图 4 caption 段落」文本替换为三子图 caption（Times New Roman 9pt bold + 宋体）。
  5. 把待办清单里残留的「（图 5）」改为「（图 4）」。

用法（项目根 /home/saw/ros2_ws/src/semphyspc/ 下）：
    python3 scripts/merge_fig45_docx.py
"""
import shutil

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.text.paragraph import Paragraph

SRC = 'docs/语义地图.docx'
BAK = 'docs/语义地图.docx.bak-20261003-fig45merge'
IMG = 'paper/figures/fig4_system_integration.png'
CAPTION = ('图 4　真实地图语义软代价导航闭环（Semantic3D domfountain_station3）：'
           '(a) 原始点云俯视；(b) 语义软代价层（连续代价：地面 0 / 植被 46 / '
           '硬刚性 150~220 / 行人 254）；(c) Nav2 规划路径（绕开行人、穿越植被）')


def para_text(el):
    return ''.join(t.text or '' for t in el.findall('.//' + qn('w:t')))


def main():
    shutil.copy2(SRC, BAK)
    doc = docx.Document(SRC)
    body = doc.element.body
    children = list(body)

    # 按 body 元素索引定位（合并前的结构，索引与探查一致）
    img4 = children[124]   # 图 4 图片段落
    cap4 = children[125]   # 图 4 caption
    text_real = children[126]  # 真实地图文字
    img5 = children[127]   # 图 5 图片段落
    cap5 = children[128]   # 图 5 caption

    assert para_text(img5).strip() == '' and img5.findall('.//' + qn('w:drawing')), 'img5 定位失败'
    assert para_text(cap5).startswith('图 5'), f'cap5 定位失败: {para_text(cap5)[:20]}'
    assert para_text(img4).strip() == '' and img4.findall('.//' + qn('w:drawing')), 'img4 定位失败'
    assert para_text(cap4).startswith('图 4'), f'cap4 定位失败: {para_text(cap4)[:20]}'

    # 1) 删除图 5 图片 + caption
    body.remove(img5)
    body.remove(cap5)

    # 2) 把真实地图文字移到图 4 图片之前（Gazebo 文字之后）
    img4.addprevious(text_real)

    # 3) 替换图 4 图片段落内的图片
    p_img = Paragraph(img4, doc)
    for r in list(p_img.runs):
        r._r.getparent().remove(r._r)
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p_img.add_run()
    run.add_picture(IMG, width=Inches(6.0))

    # 4) 替换图 4 caption 文本
    p_cap = Paragraph(cap4, doc)
    for r in list(p_cap.runs):
        r._r.getparent().remove(r._r)
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cr = p_cap.add_run(CAPTION)
    cr.font.name = 'Times New Roman'
    cr.font.size = Pt(9)
    cr.font.bold = True
    cr._element.get_or_add_rPr()
    rf = cr._element.rPr.get_or_add_rFonts()
    rf.set(qn('w:eastAsia'), '宋体')

    # 5) 待办清单「（图 5）」→「（图 4）」
    for p in doc.paragraphs:
        if '（图 5）' in p.text:
            for r in p.runs:
                if '（图 5）' in r.text:
                    r.text = r.text.replace('（图 5）', '（图 4）')

    doc.save(SRC)

    # 验证
    d2 = docx.Document(SRC)
    print(f'完成：{SRC} 已更新（备份 -> {BAK}）')
    print(f'图片数：{len(d2.inline_shapes)}')
    # 打印 §4.6 区域段落
    for p in d2.paragraphs:
        t = p.text.strip()
        if t.startswith(('语义软代价层 Gazebo', '真实地图语义导航验证', '图 4　')) or '（图 4）' in t:
            print(f'  - {t[:70]}')


if __name__ == '__main__':
    main()
