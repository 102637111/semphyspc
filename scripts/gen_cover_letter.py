#!/usr/bin/env python3
"""生成 IEEE Sensors Journal 投稿 cover letter（docs/cover_letter.docx，模板化占位）。"""
import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt

TITLE = ('Semantic Point-Cloud Segmentation towards Physical Interaction: '
         'Collision Modeling and Soft-Collision Decision for Autonomous Navigation')

body = [
    ('Dear Editor-in-Chief of IEEE Sensors Journal,', False),
    ('', False),
    ('We are pleased to submit the enclosed manuscript, "' + TITLE + '", for consideration as a regular paper '
     'in IEEE Sensors Journal.', False),
    ('', False),
    ('Contribution. The paper addresses a gap between per-point semantic segmentation and physical interaction: '
     'semantic labels alone cannot tell a planner "what happens if I hit it". We propose an end-to-end method that '
     '(i) segments point clouds with a PointMambaSeg front-end using early RGB+XYZ fusion, (ii) maps the 32 nuScenes '
     'classes to physical interaction properties (collision-safety level, mass, friction, restitution, '
     'destructibility), and (iii) turns the resulting per-point lethality into a soft-collision decision and a '
     'planner-consumable semantic soft-cost layer. The key results are: 0.3923 mIoU on nuScenes, +0.0331 mIoU from '
     'strictly paired RGB early fusion, a depth-gated 2D/3D voting mechanism that cuts the critical rate of lethal '
     'misclassification from 5.30% to 3.80%, and a Gazebo closed-loop simulation in which the planner bypasses a '
     'lethal pedestrian while driving through traversable vegetation using only the continuous soft cost.', False),
    ('', False),
    ('Fit to journal. The work integrates multimodal sensing (LiDAR + RGB), deep point-cloud segmentation, physical '
     'material modeling, and navigation safety — topics well within the scope of IEEE Sensors Journal, and of direct '
     'interest to its readership in sensor fusion and sensor-driven decision systems for autonomous vehicles.', False),
    ('', False),
    ('Statements. The manuscript is original, has not been published elsewhere, and is not under consideration by any '
     'other journal. All authors have approved the submission. The nuScenes dataset is public; code, checkpoints and '
     'supplementary results will be released upon acceptance, as stated in the manuscript.', False),
    ('', False),
    ('We thank you for your time and consideration and look forward to your response.', False),
    ('', False),
    ('Sincerely,', False),
    ('[Corresponding Author Name]', True),
    ('[Department, University, City, Country]', False),
    ('[email]', False),
]


def main():
    doc = docx.Document()
    st = doc.styles['Normal']
    st.font.name = 'Times New Roman'
    st.font.size = Pt(11)

    for text, bold in body:
        p = doc.add_paragraph()
        if text:
            run = p.add_run(text)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(11)
            run.font.bold = bold
            rf = run._element.get_or_add_rPr().get_or_add_rFonts()
            rf.set(qn('w:eastAsia'), '宋体')
        else:
            p.paragraph_format.space_after = Pt(6)

    doc.save('docs/cover_letter.docx')
    print('已生成 docs/cover_letter.docx')


if __name__ == '__main__':
    main()
