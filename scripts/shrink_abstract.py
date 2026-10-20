#!/usr/bin/env python3
"""摘要精简到 IEEE 150–250 词（双版同步）。

英文 491 词 → 206 词；中文 1196 字 → 精简稿。
整段替换摘要正文段，保留段落样式，替换后重设字号（英文 9pt / 中文 10.5pt）。

用法（项目根下）：python3 scripts/shrink_abstract.py
"""
import re
import shutil

import docx
from docx.shared import Pt
from docx.oxml.ns import qn

EN = dict(
    path='docs/paper_draft_semphyspc.docx',
    prefix='In loss-of-control scenarios',
    size=9,
    text=('In loss-of-control scenarios, autonomous vehicles should actively select non-lethal, deformable or '
          'frangible objects (e.g., bushes, traffic cones, guardrails) for soft collision to decelerate safely, '
          'while avoiding lethal or rigid targets such as pedestrians, animals, vehicles, and walls. The bottleneck '
          'is that per-point semantic labels alone are insufficient; the semantics must be translated into physical '
          'interaction properties (collision-safety level, mass, friction, restitution, destructibility). We propose '
          'an end-to-end method from semantic segmentation to physical interaction, comprising a PointMambaSeg '
          'front-end with early RGB+XYZ fusion, a 32-class semantic-to-physical-material mapping library, '
          'impact-consequence calibration, a soft-collision decision module, and navigation-system export for ROS2. '
          'On the nuScenes benchmark, the segmentation front-end reaches 0.3923 mIoU, with RGB early fusion '
          'contributing +0.0331 mIoU under a strictly paired ablation. Because rare lethal classes (pedestrians, '
          'animals) are nearly wiped out by the long-tail distribution, 2D detection with depth-gated voting is '
          'introduced, pushing the critical_rate (lethal misclassified as collidable) from 5.30% down to 3.80%, with '
          'a GT-oracle upper bound of 0.92%. The soft-collision decision is exported as a semantic soft-cost layer, '
          'and a Gazebo closed-loop simulation verifies that the planner, using only this continuous cost, '
          'automatically bypasses a lethal pedestrian while driving through traversable vegetation. The method '
          'provides a practical path for navigation-safety collision.'),
)

ZH = dict(
    path='docs/语义地图.docx',
    prefix='自动驾驶车辆在失控',
    size=10.5,
    text=('自动驾驶车辆在失控等紧急场景下需要主动选择非致命、可变形或可破坏的物体（灌木、交通锥、护栏等）进行软碰撞以安全减速，'
          '同时避免撞击行人、动物、车辆与墙体等致命或刚性目标。其瓶颈在于：仅得到逐点语义类别还不够，必须将其转化为物理交互属性'
          '（碰撞安全等级、质量、摩擦、恢复系数、可破坏性）。针对这一问题，本文提出一套从语义分割到物理交互的端到端方法，'
          '包括基于 RGB+XYZ 早期融合的 PointMambaSeg 分割前置、32 类语义→物理材质映射库、撞击后果标定、软碰撞决策模块'
          '与面向 ROS2 的导航系统导出。在 nuScenes 数据集上，分割前置达到 0.3923 mIoU，严格配对消融表明 RGB 早期融合贡献 '
          '+0.0331 mIoU；针对长尾分布下罕见致命类（行人、动物）近乎全灭的问题，引入 2D 检测与深度门控投票，把致命类被误判为'
          '可撞类的 critical_rate 从 5.30% 压至 3.80%（GT oracle 上界 0.92%）；软碰撞决策落地为语义软代价层，并在 Gazebo '
          '物理仿真闭环中验证了规划器仅凭连续软代价即可绕开致命行人、穿越可穿越植被。该方法为导航安全碰撞提供了一条可落地的技术路径。'),
)


def set_text(p, text, size):
    for r in list(p.runs):
        r._r.getparent().remove(r._r)
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(size)
    rf = run._element.get_or_add_rPr().get_or_add_rFonts()
    rf.set(qn('w:eastAsia'), '宋体')


def main():
    for cfg in (EN, ZH):
        shutil.copy2(cfg['path'], cfg['path'] + '.bak-abstract-shrink')
        doc = docx.Document(cfg['path'])
        done = False
        for p in doc.paragraphs:
            if p.text.strip().startswith(cfg['prefix']):
                set_text(p, cfg['text'], cfg['size'])
                done = True
                break
        assert done, f"{cfg['path']}: 摘要段未找到"
        doc.save(cfg['path'])

        d2 = docx.Document(cfg['path'])
        for p in d2.paragraphs:
            if p.text.strip().startswith(cfg['prefix']):
                t = p.text.strip()
                words = re.findall(r"[A-Za-z0-9.]+(?:-[A-Za-z0-9.]+)*", t)
                print(f"{cfg['path']}: 摘要已精简，词数={len(words)}，字数={len(t)}，含5.30%→3.80%={'5.30%' in t and '3.80%' in t}")
                break


if __name__ == '__main__':
    main()
