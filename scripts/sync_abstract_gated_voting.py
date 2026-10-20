#!/usr/bin/env python3
"""在摘要内补一句门控投票成果（致命类误判 5.30%→3.80%），与 §4.2.3、结论呼应。

插入位置：级联实例分割/置信度门控成果之后、撞击后果标定之前。
锚点必须完整落在单个 run 内，否则报错。

用法（项目根下）：python3 scripts/sync_abstract_gated_voting.py
"""
import shutil

import docx

JOBS = [
    dict(
        path='docs/语义地图.docx',
        para_prefix='自动驾驶车辆在失控',
        anchor='并对撞击后果做了闭合形式标定',
        insert=('针对罕见类近乎全灭遗留的致命类安全误判，2D 检测与深度门控投票把致命类被误判为可撞类的 '
                'critical_rate 从 5.30% 压至 3.80%（GT oracle 上界 0.92%）；'),
    ),
    dict(
        path='docs/paper_draft_semphyspc.docx',
        para_prefix='In loss-of-control scenarios',
        anchor='we further calibrate the impact consequences',
        insert=('against the lethal-class safety misclassification left by the near-total failure of rare classes, '
                '2D detection with depth-gated voting pushes the critical_rate (lethal misclassified as collidable) '
                'from 5.30% down to 3.80%, with a GT-oracle upper bound of 0.92%; '),
    ),
]


def main():
    for job in JOBS:
        shutil.copy2(job['path'], job['path'] + '.bak-abstract')
        doc = docx.Document(job['path'])
        done = False
        for p in doc.paragraphs:
            if p.text.strip().startswith(job['para_prefix']):
                for r in p.runs:
                    if job['anchor'] in r.text:
                        r.text = r.text.replace(job['anchor'], job['insert'] + job['anchor'], 1)
                        done = True
                        break
                break
        assert done, f"{job['path']}: 摘要锚点未在单个 run 内找到"
        doc.save(job['path'])

        # 验证
        d2 = docx.Document(job['path'])
        for p in d2.paragraphs:
            if p.text.strip().startswith(job['para_prefix']):
                assert '3.80%' in p.text and '5.30%' in p.text, '摘要同步缺失'
                print(f"完成：{job['path']} 摘要已补门控投票句")
                break


if __name__ == '__main__':
    main()
