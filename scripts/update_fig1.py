#!/usr/bin/env python3
"""更新 Fig.1：替换 pipeline 图（加入门控投票）+ 更新 caption + 修正显示比例。

- 用 cairosvg 把 fig1_pipeline.svg 渲染到 2400px 宽（≥300 DPI）
- 替换 docx 里 image1.png（rId6）的字节
- 更新 drawing 的 ext cy 匹配新图比例
- 更新 Fig.1 caption 对齐新主线

用法（项目根下）：python3 scripts/update_fig1.py
"""
import shutil

import cairosvg
import docx
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.text.paragraph import Paragraph

SRC = 'docs/paper_draft_semphyspc.docx'
BAK = SRC + '.bak-fig1'
SVG = 'paper/figures/fig1_pipeline.svg'
PNG = 'paper/figures/fig1_pipeline.png'

NEW_CAPTION = (
    'Fig. 1. The SemPhysPC overall pipeline: the perception front-end combines PointMambaSeg segmentation with '
    'depth-gated voting, and the semantic-to-physical mapping then drives the downstream tasks — '
    'impact-consequence calibration, soft-collision decision and cross-frame tracking.'
)


def text(el):
    ts = el.findall('.//' + qn('w:t'))
    return ''.join(t.text or '' for t in ts)


def set_para_text(p, txt, size=12, bold=False):
    for r in list(p.runs):
        r._r.getparent().remove(r._r)
    run = p.add_run(txt)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(size)
    run.font.bold = bold
    rf = run._element.get_or_add_rPr().get_or_add_rFonts()
    rf.set(qn('w:eastAsia'), '宋体')
    return run


def main():
    shutil.copy2(SRC, BAK)
    doc = docx.Document(SRC)

    # 1. 渲染新图到 2400 宽
    png_bytes = cairosvg.svg2png(url=SVG, output_width=2400)
    with open(PNG, 'wb') as f:
        f.write(png_bytes)
    print(f'渲染新 Fig.1: {len(png_bytes)} bytes')

    import io
    from PIL import Image
    _w, _h = Image.open(io.BytesIO(png_bytes)).size
    ratio = _h / _w
    print(f'渲染尺寸: {_w}x{_h}, 比例 {ratio:.4f}')

    # 2. 替换 docx 里 image1.png 的 blob，并记下其 rId
    target_rid = None
    for rid, part in doc.part.related_parts.items():
        if str(part.partname) == '/word/media/image1.png':
            part._blob = png_bytes
            target_rid = rid
            break
    assert target_rid is not None, '未找到 image1.png'
    print(f'已替换 image1.png 字节 (rId={target_rid})')

    # 3. 修正显示比例（ext cy）
    # 用渲染后 PNG 的实际比例；保持 cx 不变，改 cy
    n_ext = 0
    for el in doc.element.body:
        # 定位含 image1.png 的 drawing
        blips = el.findall('.//' + qn('a:blip'))
        for b in blips:
            if b.get(qn('r:embed')) == target_rid:
                ext = el.find('.//' + qn('wp:extent'))
                cx = int(ext.get('cx'))
                ext.set('cy', str(int(cx * ratio)))
                print(f'更新 ext cy -> {int(cx * ratio)} EMU')
                n_ext += 1
                break
    assert n_ext >= 1, '未找到 Fig.1 drawing ext'

    # 4. 更新 caption
    cap = None
    for el in doc.element.body:
        if el.tag == qn('w:p') and text(el).strip().startswith('Fig. 1. The SemPhysPC overall pipeline'):
            cap = el
            break
    assert cap is not None, 'Fig.1 caption 未找到'
    set_para_text(Paragraph(cap, None), NEW_CAPTION)
    print('已更新 Fig.1 caption')

    doc.save(SRC)

    # ---- 验证 ----
    d2 = docx.Document(SRC)
    # 图片尺寸
    import io
    from PIL import Image
    for rid, part in d2.part.related_parts.items():
        if str(part.partname) == '/word/media/image1.png':
            im = Image.open(io.BytesIO(part.blob))
            print(f'验证 image1.png 尺寸: {im.size}')
    joined = '\n'.join(p.text for p in d2.paragraphs)
    print('caption 含门控投票:', 'OK' if 'depth-gated voting' in joined else 'FAIL')


if __name__ == '__main__':
    main()
