#!/usr/bin/env python3
"""审查修复（3 处）：

1. 门控投票正文段位置错误（在标题之前）→ 移到标题之后
2. 3.1 Overview pipeline 缺门控投票环节 → 加入
3. 3.2 扩充段与 [26] 重复（FPS/num_group/缩放瓶颈）→ 精简为形式化符号 + 线性复杂度 + KNN 上采样

用法（项目根下）：python3 scripts/fix_review_issues.py
"""
import shutil

import docx
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.text.paragraph import Paragraph

SRC = 'docs/paper_draft_semphyspc.docx'
BAK = SRC + '.bak-fix'

NEW_21 = (
    'The pipeline consists of a perception front-end (segmentation and depth-gated voting) and a main body '
    '(the physical-interaction pipeline):'
)

NEW_22 = (
    'LiDAR + camera RGB ──► ① PointMambaSeg (early RGB+XYZ fusion, 32 classes) → '
    '② depth-gated 2D/3D voting (recovers rare lethal classes) → ③ semantic-to-physical '
    'material mapping → ④ instance-level clustering (DBSCAN, thing/stuff partition) → '
    '⑤ impact-consequence calibration / ⑥ soft-collision decision / ⑦ cross-frame tracking '
    '→ downstream (navigation ROS2 nodes, map-level fusion fused_map.npz, extensible game export).'
)

NEW_30 = (
    'Formally, given a LiDAR point cloud X ∈ R^{N×3} and the camera RGB projected onto the same '
    'points C ∈ R^{N×3}, the front-end concatenates them into a six-channel tensor [X, C]. The '
    'non-hierarchical encoder applies Mamba state-space blocks with linear complexity in the sequence length, '
    'so colour and geometry are jointly modeled as a single unified representation before any grouping; a '
    'KNN-upsampling decoder then propagates the grouped features back to the full resolution to recover '
    'per-point semantics.'
)


def text(el):
    ts = el.findall('.//' + qn('w:t'))
    return ''.join(t.text or '' for t in ts)


def set_para_text(p, txt, size=12, bold=False):
    for r in list(p.runs):
        r._r.getparent().remove(r._r)
    run = p.add_run(txt)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(size)
    run.font.bold = bold
    rf = run._element.get_or_add_rPr().get_or_add_rFonts()
    rf.set(qn('w:eastAsia'), '宋体')
    return run


def find_body_para(doc, prefix):
    for el in doc.element.body:
        if el.tag == qn('w:p') and text(el).strip().startswith(prefix):
            return el
    return None


def main():
    shutil.copy2(SRC, BAK)
    doc = docx.Document(SRC)

    # ---- 1. 门控投票正文移到标题之后 ----
    el_title = find_body_para(doc, '3.3 Depth-Gated 2D/3D Voting')
    assert el_title is not None, '门控投票标题未找到'
    el_body = el_title.getprevious()
    # 正文应是标题前一个段落
    assert el_body is not None and text(el_body).strip().startswith('Rare thing classes'), \
        '门控投票正文定位失败'
    el_title.addnext(el_body)   # lxml 会先把 el_body 从原位置摘下再插到标题后

    # ---- 2. Overview pipeline 加入门控投票 ----
    el_21 = find_body_para(doc, 'The pipeline consists of two stages')
    el_22 = find_body_para(doc, 'LiDAR + camera RGB')
    assert el_21 is not None and el_22 is not None, 'Overview 定位失败'
    set_para_text(Paragraph(el_21, None), NEW_21)
    set_para_text(Paragraph(el_22, None), NEW_22)

    # ---- 3. 3.2 扩充段精简 ----
    el_30 = find_body_para(doc, 'Formally, given a LiDAR point cloud')
    assert el_30 is not None, '3.2 扩充段未找到'
    set_para_text(Paragraph(el_30, None), NEW_30)

    doc.save(SRC)

    # ---- 验证 ----
    d2 = docx.Document(SRC)
    body = d2.element.body
    def t(el):
        return ''.join(x.text or '' for x in el.findall('.//' + qn('w:t')))
    # 1. 顺序：标题后紧跟正文
    ok_order = False
    for i, el in enumerate(body):
        if t(el).strip().startswith('3.3 Depth-Gated'):
            nxt = el.getnext()
            ok_order = nxt is not None and t(nxt).strip().startswith('Rare thing classes')
            break
    print('1. 门控投票标题后紧跟正文:', 'OK' if ok_order else 'FAIL')
    # 2. pipeline 含门控投票
    joined = '\n'.join(p.text for p in d2.paragraphs)
    print('2. pipeline 含门控投票:', 'OK' if 'depth-gated 2D/3D voting (recovers rare lethal classes)' in joined else 'FAIL')
    print('   pipeline 步骤 7 个:', 'OK' if '⑦ cross-frame tracking' in joined else 'FAIL')
    # 3. 重复内容已删
    print('3. FPS/num_group 重复已删:', 'OK' if 'farthest-point sampling (FPS) selects num_group' not in joined else 'FAIL')


if __name__ == '__main__':
    main()
