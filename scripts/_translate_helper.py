#!/usr/bin/env python3
"""翻译辅助函数：把 docx 段落/单元格文本替换为英文译文（保留段落样式与 bold 状态）。"""
from docx.oxml.ns import qn


def set_text(p, text, bold=None):
    """清空段落 runs，写入纯文本（Times New Roman），保留段落样式与 bold 状态。"""
    if bold is None:
        bold = any(r.bold for r in p.runs if (r.text or '').strip())
    for r in list(p.runs):
        r._r.getparent().remove(r._r)
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    if bold:
        run.font.bold = True
    return run


def apply_prefix(doc, trans):
    """trans: {唯一前缀: 英文译文}。匹配前缀的段落整段替换，返回替换条数。"""
    n = 0
    for p in doc.paragraphs:
        t = p.text.strip()
        for pre, tr in trans.items():
            if t.startswith(pre):
                set_text(p, tr)
                n += 1
                break
    return n


def delete_prefix(doc, prefixes):
    """删除文本以任一前缀开头的段落，返回删除条数。"""
    removed = 0
    for p in list(doc.paragraphs):
        t = p.text.strip()
        if any(t.startswith(pre) for pre in prefixes):
            p._p.getparent().remove(p._p)
            removed += 1
    return removed


def delete_exact(doc, texts):
    """删除文本完全等于 texts 中任一值的段落，返回删除条数。"""
    removed = 0
    for p in list(doc.paragraphs):
        if p.text.strip() in texts:
            p._p.getparent().remove(p._p)
            removed += 1
    return removed


def set_cell_text(cell, text, bold=None):
    """替换单元格文本（用单元格首段）。"""
    set_text(cell.paragraphs[0], text, bold=bold)
