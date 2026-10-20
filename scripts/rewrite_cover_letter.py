#!/usr/bin/env python3
"""重写 cover letter 针对 MDPI Remote Sensing 特刊
「LiDAR Technology for Autonomous Navigation and Mapping」。

覆盖 docs/cover_letter.docx（先备份 .bak）。

用法（项目根下）：python3 scripts/rewrite_cover_letter.py
"""
import shutil

import docx
from docx.shared import Pt

SRC = 'docs/cover_letter.docx'
BAK = SRC + '.bak'

TITLE = ('Semantic Point-Cloud Segmentation with RGB–LiDAR Early Fusion '
         'and Depth-Gated Voting for Autonomous Navigation')

PARAS = [
    ('Dear Editors of the Special Issue “LiDAR Technology for Autonomous Navigation and '
     'Mapping”,'),
    ('We are pleased to submit the enclosed manuscript, “' + TITLE + '”, for consideration '
     'for publication in the Special Issue “LiDAR Technology for Autonomous Navigation and '
     'Mapping” of Remote Sensing.'),
    ('The manuscript presents SemPhysPC, a semantic point-cloud segmentation method oriented to '
     'autonomous navigation. Its main contributions are: (i) a PointMambaSeg front-end that fuses '
     'camera RGB with LiDAR XYZ through an early-fusion six-channel input, with a strictly paired '
     'ablation quantifying the RGB contribution (+0.0331 mIoU, 0.3923 mIoU on nuScenes); (ii) a '
     'depth-gated 2D/3D voting mechanism that lifts high-confidence 2D detections into the 3D labels '
     'to recover the rare but safety-critical lethal classes, lowering the critical misclassification '
     'rate from 5.30% to 3.80% (GT-oracle upper bound 0.92%); and (iii) a semantic-to-physical '
     'attribute mapping exported as a continuous semantic soft-cost layer for a ROS2 navigation '
     'stack. A Gazebo closed-loop simulation and a Semantic3D real-map experiment verify that the '
     'planner, using only this soft cost, bypasses a lethal pedestrian while driving through '
     'traversable vegetation.'),
    ('The work sits squarely within the scope of the Special Issue: it combines LiDAR point-cloud '
     'segmentation, multimodal RGB–LiDAR fusion, and navigation-oriented physical reasoning, and '
     'it is validated both in a closed-loop simulator and on a real LiDAR-derived map. We believe it '
     'will interest the Special Issue readership working on LiDAR perception, semantic mapping, and '
     'safe navigation.'),
    ('This manuscript is original, has not been published previously, and is not under consideration '
     'for publication elsewhere. All authors have read and approved the manuscript and agree to its '
     'submission to Remote Sensing. The nuScenes and Semantic3D datasets are public; the mapping '
     'library, trained models, and source code will be released as described in the Data Availability '
     'Statement.'),
    ('We thank you for your time and consideration and look forward to hearing from you.'),
    ('Sincerely,'),
    ('[Corresponding Author Name]'),
    ('[Department, University, City, Country]'),
    ('[email]'),
]


def main():
    shutil.copy2(SRC, BAK)
    doc = docx.Document(SRC)
    # 清空现有段落
    for p in list(doc.paragraphs):
        p._element.getparent().remove(p._element)
    # 写入新内容
    for txt in PARAS:
        p = doc.add_paragraph(txt)
        p.paragraph_format.line_spacing = 1.5
        for r in p.runs:
            r.font.name = 'Times New Roman'
            r.font.size = Pt(12)
    doc.save(SRC)

    # 验证
    d2 = docx.Document(SRC)
    joined = ' '.join(p.text for p in d2.paragraphs)
    print('段数:', len(d2.paragraphs))
    print('特刊名在:', 'LiDAR Technology for Autonomous Navigation and Mapping' in joined)
    print('标题在:', 'Depth-Gated Voting' in joined)
    print('关键词在:', all(s in joined for s in ['0.3923', '3.80%', '0.92%', 'soft-cost']))
    print('--- 前 3 段 ---')
    for p in d2.paragraphs[:3]:
        print(' ', p.text[:80])


if __name__ == '__main__':
    main()
