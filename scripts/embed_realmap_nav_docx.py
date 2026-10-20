#!/usr/bin/env python3
"""把「真实地图（Semantic3D）语义导航验证」填补进权威载体 docs/语义地图.docx。

在 §4.7「语义软代价层 Gazebo 闭环」（图 4）之后插入：
  1. 文字 bullet「真实地图语义导航验证」（复用 Gazebo bullet 格式 + **bold** 解析）
  2. 居中图片段落（路径规划地图.png，宽 6.00in，与图 4 全宽图一致）
  3. 居中 caption 段落（图 5　…，Times New Roman 9pt bold + 宋体中文回退）
  4. 待办清单末尾插入本条目

用法（项目根 /home/saw/ros2_ws/src/semphyspc/ 下）：
    python3 scripts/embed_realmap_nav_docx.py
"""
import copy
import re
import shutil

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.text.paragraph import Paragraph

SRC = 'docs/语义地图.docx'
BAK = 'docs/语义地图.docx.bak-20261002-realmap'
IMG = 'paper/figures/effect/路径规划地图.png'
CAPTION = '图 5　真实地图（Semantic3D domfountain_station3）语义软代价路径规划（绕开行人、穿越植被）'

BULLET = (
    '**真实地图语义导航验证**：将语义软代价层从 Gazebo 合成场景推广到**真实室外大场景**——'
    'Semantic3D 的 domfountain_station3（约 3520 万点）经分割前置推理后，投影为 50×50 m 的语义代价栅格'
    '（resolution 0.2 m，origin (−25, −25)）。静态地图仅将致命类（行人）标为 occupied，'
    '硬刚性类（车/墙）与植被留为 free，由语义软代价层给出连续代价'
    '（致命 254 / 硬刚性 150~220 / 植被 46 / 地面 0）。'
    'Nav2（NavfnPlanner）在 A(−13.9, 3.7) → B(12.5, −3.1) 上规划得到 296 点路径：'
    '全程不经过任何致命 cell（离最近行人 1.68 m）、不经过硬刚性 cell、穿越 4 个植被软代价 cell、'
    '98.6% 落于地面 free cell——规划器仅凭连续软代价即在真实点云地图上自动「绕开行人、穿越植被」，'
    '验证了语义软代价层对真实地图的适用性。'
)

TODO = '• [x] 真实地图（Semantic3D domfountain_station3）语义软代价导航验证已写入 §4.7（图 5）'


def find_para(doc, needle, exclude=None):
    for p in doc.paragraphs:
        if needle in p.text:
            if exclude and any(e in p.text for e in exclude):
                continue
            return p
    raise RuntimeError(f'段落未找到: {needle}')


def _set_run_font(r, bold):
    r.font.name = 'Times New Roman'
    r.font.size = Pt(10.5)
    r.font.bold = bold
    r._element.get_or_add_rPr()
    rf = r._element.rPr.get_or_add_rFonts()
    rf.set(qn('w:eastAsia'), '宋体')


def make_md_para(template_para, text):
    """deepcopy 模板段落，按 **bold** / `code` 分段填充，返回新 w:p 元素。"""
    el = copy.deepcopy(template_para._p)
    p = Paragraph(el, template_para._parent)

    bold_src = normal_src = None
    for r in p.runs:
        if r.font.bold and bold_src is None:
            bold_src = r._r
        elif not r.font.bold and normal_src is None:
            normal_src = r._r
    for r in list(p.runs):
        r._r.getparent().remove(r._r)

    for tok in re.split(r'(\*\*[^*]+\*\*|`[^`]*`)', text):
        if not tok:
            continue
        bold = False
        if tok.startswith('**') and tok.endswith('**'):
            bold = True
            tok = tok[2:-2]
        elif tok.startswith('`') and tok.endswith('`'):
            tok = tok[1:-1]
        src = (bold_src if bold else normal_src) or bold_src or normal_src
        new_r = copy.deepcopy(src)
        t_el = new_r.find(qn('w:t'))
        if t_el is None:
            t_el = new_r.makeelement(qn('w:t'), {})
            new_r.append(t_el)
        t_el.text = tok
        rPr = new_r.find(qn('w:rPr'))
        if rPr is None:
            rPr = new_r.makeelement(qn('w:rPr'), {})
            new_r.insert(0, rPr)
        b = rPr.find(qn('w:b'))
        if bold and b is None:
            rPr.append(rPr.makeelement(qn('w:b'), {}))
        elif not bold and b is not None:
            rPr.remove(b)
        p._p.append(new_r)
    return el


def new_para_like(template, text):
    el = copy.deepcopy(template._p)
    p = Paragraph(el, template._parent)
    for r in p.runs:
        r.text = ''
    if p.runs:
        p.runs[0].text = text
    else:
        p.add_run(text)
    return el


def main():
    shutil.copy2(SRC, BAK)
    doc = docx.Document(SRC)

    # 图 1 的图片段落与 caption 段落作为格式模板
    p_fig1_img = p_fig1_cap = None
    for i, p in enumerate(doc.paragraphs):
        if p.text.startswith('图 1　') and p_fig1_cap is None:
            p_fig1_cap = p
            p_fig1_img = doc.paragraphs[i - 1]
    assert p_fig1_img is not None and p_fig1_cap is not None

    # bullet 模板 = Gazebo 闭环 bullet（在文档中先于图 4 caption，天然先匹配）
    # 锚点 = 图 4 caption（在其后插入）
    p_bullet_tpl = find_para(doc, '语义软代价层 Gazebo 闭环')
    anchor = find_para(doc, '图 4　语义软代价层 Gazebo 闭环轨迹')

    # 1) 文字 bullet
    bullet_el = make_md_para(p_bullet_tpl, BULLET)

    # 2) 图片段落（deepcopy 图1图片段落，清空旧图，加路径规划地图）
    img_el = copy.deepcopy(p_fig1_img._p)
    img_p = Paragraph(img_el, p_fig1_img._parent)
    for r in list(img_p.runs):
        r._r.getparent().remove(r._r)
    img_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    img_p.add_run().add_picture(IMG, width=Inches(6.0))

    # 3) caption 段落（deepcopy 图1 caption，替换文本）
    cap_el = copy.deepcopy(p_fig1_cap._p)
    cap_p = Paragraph(cap_el, p_fig1_cap._parent)
    for r in list(cap_p.runs):
        r._r.getparent().remove(r._r)
    cap_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cr = cap_p.add_run(CAPTION)
    cr.font.name = 'Times New Roman'
    cr.font.size = Pt(9)
    cr.font.bold = True
    cr._element.get_or_add_rPr()
    rf = cr._element.rPr.get_or_add_rFonts()
    rf.set(qn('w:eastAsia'), '宋体')

    # 插入：图 4 caption 之后 → bullet → 图片 → caption
    anchor._p.addnext(bullet_el)
    bullet_el.addnext(img_el)
    img_el.addnext(cap_el)

    # 4) 待办清单末尾插入条目
    p_todo = find_para(doc, '阶段 3 语义软代价层 Gazebo 闭环')
    p_todo._p.addnext(new_para_like(p_todo, TODO))

    doc.save(SRC)

    d2 = docx.Document(SRC)
    n_img = len(d2.inline_shapes)
    last = d2.inline_shapes[-1]
    print(f'完成：{SRC} 已更新（备份 -> {BAK}）')
    print(f'图片数：{n_img}（上一版图 4 后应为 9）')
    print(f'最后一张图宽 {last.width/914400:.2f}in x 高 {last.height/914400:.2f}in')
    # 校验图 5 caption 与 bullet 落位
    for p in d2.paragraphs:
        if p.text.startswith('图 5　') or '真实地图语义导航验证' in p.text:
            print('  ✓', p.text[:60])


if __name__ == '__main__':
    main()
